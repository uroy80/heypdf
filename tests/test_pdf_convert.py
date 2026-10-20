"""heypdf light conversions: images → PDF, PDF → images / Word / Excel / Markdown (our own converter)."""
from __future__ import annotations

import io
import zipfile

import pymupdf
import pytest
from PIL import Image

from app.pdf.convert import to_number
from app.pdf.errors import PdfError
from tests.fixtures.pdf.factory import pdfs, run_tool  # noqa: F401  (session fixture)


def test_images_to_pdf_layout_and_jpeg_passthrough(pdfs, tmp_path):
    ctx = run_tool("pdf.from_images", tmp_path, [(pdfs.jpeg, "image"), (pdfs.png_alpha, "image")],
                   {"page_size": "a4", "orientation": "portrait", "margin": 20})
    doc = pymupdf.open(ctx.outputs[0][0])
    assert doc.page_count == 2
    assert doc[0].rect.width == pytest.approx(595.3, abs=1) and doc[0].rect.height == pytest.approx(841.9, abs=1)
    xref = doc[0].get_images()[0][0]
    assert doc.extract_image(xref)["image"] == pdfs.jpeg.read_bytes()  # lossless: the JPEG is embedded as is
    separate = run_tool("pdf.from_images", tmp_path, [(pdfs.jpeg, "image"), (pdfs.png_alpha, "image")],
                        {"merge": False, "page_size": "fit"})
    assert len(separate.outputs) == 2
    assert pymupdf.open(separate.outputs[0][0])[0].rect.width == pytest.approx(640 * 72 / 96, rel=0.35)


def _image_inputs(tmp_path):
    """HEIC (iPhone), SVG (logo) and a 16-bit greyscale PNG gradient (scanner)."""
    import numpy as np

    from tests.conftest import make_heic_bytes

    heic = tmp_path / "p.heic"
    heic.write_bytes(make_heic_bytes((300, 200)))
    svg = tmp_path / "logo.svg"
    svg.write_text('<svg xmlns="http://www.w3.org/2000/svg" width="200" height="100">'
                   '<rect width="200" height="100" fill="#cc0000"/></svg>')
    g16 = tmp_path / "g16.png"
    arr = np.tile(np.linspace(512, 65280, 256).astype(np.uint16), (128, 1))
    Image.fromarray(arr).save(g16)  # uint16 → mode I;16
    assert Image.open(g16).mode == "I;16"
    return heic, svg, g16


def _page_pixels(path, page_no=0):
    import numpy as np

    pix = pymupdf.open(path)[page_no].get_pixmap(dpi=36, alpha=False)
    return np.frombuffer(pix.samples, np.uint8).reshape(pix.h, pix.w, pix.n)


def test_image_formats_reach_every_pdf_tool(pdfs, tmp_path):
    """HEIC and SVG are read (not 415) and a 16-bit PNG is not a blank white page, in JPG → PDF,
    Scan, Watermark, Sign and Edit (they all decode through the shared image loader)."""
    import base64

    heic, svg, g16 = _image_inputs(tmp_path)
    ctx = run_tool("pdf.from_images", tmp_path, [(heic, "image"), (svg, "image"), (g16, "image")],
                   {"page_size": "fit"})
    doc = pymupdf.open(ctx.outputs[0][0])
    assert doc.page_count == 3
    red = _page_pixels(ctx.outputs[0][0], 1)
    assert red[..., 0].mean() > 150 and red[..., 1].mean() < 80  # the SVG's red rect, not a blank page
    grey = _page_pixels(ctx.outputs[0][0], 2)
    assert grey.min() < 30 and grey.max() > 225 and 90 < grey.mean() < 170  # the gradient survives
    scan = run_tool("pdf.scan", tmp_path, [(g16, "image"), (heic, "image"), (svg, "image")],
                    {"detect": False, "enhance": "none", "page_size": "fit"})
    page = _page_pixels(scan.outputs[0][0], 0)
    assert page.min() < 40 and page.mean() < 200
    for img in (heic, svg):
        wm = run_tool("pdf.watermark", tmp_path, [pdfs.report], {"kind": "image", "image_id": "$img", "pages": "1"},
                      extras={"$img": (img, "image")})
        assert len(pymupdf.open(wm.outputs[0][0])[0].get_images()) > len(pymupdf.open(pdfs.report)[0].get_images())
        place = [{"x": 0.1, "y": 0.1, "w": 0.3, "h": 0.1, "pages": "1"}]
        signed = run_tool("pdf.sign", tmp_path, [pdfs.report],
                          {"style": "upload", "image_id": "$img", "placements": place},
                          extras={"$img": (img, "image")})
        assert pymupdf.open(signed.outputs[0][0])[0].get_image_info()
    data = "data:image/png;base64," + base64.b64encode(g16.read_bytes()).decode()
    edited = run_tool("pdf.edit", tmp_path, [pdfs.report], {"items": [
        {"type": "image", "page": 1, "x": 0.1, "y": 0.1, "w": 0.4, "h": 0.2, "image_data": data}]})
    infos = [i for i in pymupdf.open(edited.outputs[0][0])[0].get_images(full=True)]
    assert len(infos) > len(pymupdf.open(pdfs.report)[0].get_images())


def test_images_to_pdf_rejects_non_images(pdfs, tmp_path):
    with pytest.raises(PdfError):
        run_tool("pdf.from_images", tmp_path, [(pdfs.not_pdf, "image")], {})


def test_pdf_to_images_pages_respect_rotation(pdfs, tmp_path):
    ctx = run_tool("pdf.to_images", tmp_path, [pdfs.geometry], {"format": "png", "dpi": 72})
    sizes = [Image.open(p).size for p, _n, _m in ctx.outputs]
    assert sizes[0][0] < sizes[0][1] and sizes[1][0] > sizes[1][1]  # page 2 is displayed landscape
    assert [n for _p, n, _m in ctx.outputs][0] == "geometry_page-0001.png"


def test_pdf_to_images_extract_embedded(pdfs, tmp_path):
    ctx = run_tool("pdf.to_images", tmp_path, [pdfs.report], {"mode": "extract"})
    assert len(ctx.outputs) == 1 and ctx.outputs[0][2] == "image/jpeg"
    with pytest.raises(PdfError, match="No images"):
        run_tool("pdf.to_images", tmp_path, [pdfs.geometry], {"mode": "extract"})


def test_pdf_to_word_layout_and_text_modes(pdfs, tmp_path):
    import docx

    ctx = run_tool("pdf.to_word", tmp_path, [pdfs.report], {})
    d = docx.Document(ctx.outputs[0][0])
    body = "\n".join(p.text for p in d.paragraphs)
    assert "Quarterly results" in body and ctx.data["mode"] == "layout"
    assert len(d.tables) >= 1
    # the OCR'd (invisible text) page goes through the text route automatically
    scanned = run_tool("pdf.to_word", tmp_path, [pdfs.scanned], {"pages": "2"})
    assert scanned.data["mode"] == "text"
    assert "quarterly report" in "\n".join(p.text for p in docx.Document(scanned.outputs[0][0]).paragraphs)
    # no stale template statistics ("1 page", "Microsoft Macintosh Word"): the upload probe shows no page count
    from app.core.files import sniff

    for out in (ctx.outputs[0][0], scanned.outputs[0][0]):
        with zipfile.ZipFile(out) as z:
            app_xml = z.read("docProps/app.xml")
        assert b"<Pages>" not in app_xml and b"Microsoft" not in app_xml and b"heypdf" in app_xml
        assert sniff(out).pages is None


def _docx_text(path) -> str:
    with zipfile.ZipFile(path) as z:
        xml = z.read("word/document.xml").decode("utf-8")
    import re

    return "".join(re.findall(r"<w:t(?:\s[^>]*)?>([^<]*)</w:t>", xml))


def test_pdf_to_word_keeps_pages_with_rotate_90_and_180(tmp_path):
    """Pdf2docx silently dropped pages whose /Rotate is 90 or 180 (content horizontal on the
    unrotated page — e.g. after Rotate PDF); every page marker must reach the DOCX in layout mode."""
    src = tmp_path / "rotated.pdf"
    doc = pymupdf.open()
    for i, rot in enumerate((0, 90, 180, 90, 270, 0), start=1):
        page = doc.new_page()
        page.insert_text((72, 100), f"MARKERPAGE{i} quarterly figures", fontsize=14)
        page.insert_text((72, 130), f"Region totals for section {i}", fontsize=11)
        page.set_rotation(rot)
    doc.save(src)
    for mode in ("auto", "layout"):
        ctx = run_tool("pdf.to_word", tmp_path, [src], {"mode": mode})
        text = _docx_text(ctx.outputs[0][0])
        assert all(f"MARKERPAGE{i}" in text for i in range(1, 7)), (mode, text[:300])
        assert ctx.data["mode"] == "layout" and not ctx.warnings
    assert [p.rotation for p in pymupdf.open(src)] == [0, 90, 180, 90, 270, 0]  # the input is untouched


def test_pdf_to_word_falls_back_to_text_when_a_page_comes_out_empty(tmp_path, monkeypatch):
    from app.pdf import convert

    src = tmp_path / "plain.pdf"
    doc = pymupdf.open()
    for i in (1, 2):
        doc.new_page().insert_text((72, 100), f"SAFETYNET{i} content", fontsize=12)
    doc.save(src)
    out = tmp_path / "x.docx"
    convert._docx_from_text(pymupdf.open(src), [1], out)  # a DOCX holding page 1 only
    assert convert._pages_missing_from_docx(pymupdf.open(src), [1, 2], out) == [2]
    assert convert._pages_missing_from_docx(pymupdf.open(src), [1], out) == []


_PROSE = ("The committee met on Tuesday to discuss the budget for the coming year and agreed that the "
          "new library wing should be finished before the monsoon arrives in the region next summer. "
          "Members raised questions about the contractor, the cost of materials and the timeline. ") * 3


def test_pdf_to_excel_auto_does_not_turn_prose_into_tables(tmp_path):
    """Without ruled tables, 'auto' used to accept any whitespace 'table' — words cut in half."""
    one = tmp_path / "prose.pdf"
    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_textbox(pymupdf.Rect(72, 72, 540, 700), _PROSE * 2, fontsize=11)
    doc.save(one)
    two = tmp_path / "twocol.pdf"
    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_text((72, 60), "Two Column Article Title", fontsize=16)
    page.insert_textbox(pymupdf.Rect(72, 80, 296, 760), _PROSE * 2, fontsize=10)
    page.insert_textbox(pymupdf.Rect(316, 80, 540, 760), _PROSE * 2, fontsize=10)
    doc.save(two)
    for src in (one, two):
        with pytest.raises(PdfError, match="No tables were found"):
            run_tool("pdf.to_excel", tmp_path, [src], {})


def test_pdf_to_excel_auto_still_finds_borderless_statements(tmp_path):
    import openpyxl

    src = tmp_path / "statement.pdf"
    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_text((72, 60), "Account statement", fontsize=14)
    rows = [("Date", "Description", "Amount", "Balance")] + [
        (f"2026-10-{d:02d}", desc, amt, bal) for d, desc, amt, bal in (
            (1, "Opening balance", "0.00", "1,000.00"), (3, "Grocery store", "-45.20", "954.80"),
            (7, "Salary October", "2,500.00", "3,454.80"), (9, "Electricity bill", "-120.00", "3,334.80"),
            (15, "Coffee shop", "-4.50", "3,330.30"), (21, "Book store", "-18.99", "3,311.31"))]
    for i, row in enumerate(rows):
        for x, text in zip((72, 160, 330, 430), row, strict=True):
            page.insert_text((x, 100 + 18 * i), text, fontsize=10)
    doc.save(src)
    ctx = run_tool("pdf.to_excel", tmp_path, [src], {})
    assert ctx.data["tables"] >= 1
    ws = openpyxl.load_workbook(ctx.outputs[0][0]).worksheets[0]
    values = [[c.value for c in row if c.value is not None] for row in ws.iter_rows()]
    flat = [v for row in values for v in row]
    assert "Salary October" in flat and "Grocery store" in flat and 3454.8 in flat


def test_pdf_to_excel_tables_and_numbers(pdfs, tmp_path):
    import openpyxl

    ctx = run_tool("pdf.to_excel", tmp_path, [pdfs.report], {})
    wb = openpyxl.load_workbook(ctx.outputs[0][0])
    assert len(wb.sheetnames) == ctx.data["tables"] == 3
    ws = wb[wb.sheetnames[0]]
    assert ws["A1"].value == "Region" and ws["B2"].value == 107 and isinstance(ws["B2"].value, int)
    with pytest.raises(PdfError, match="No tables"):
        run_tool("pdf.to_excel", tmp_path, [pdfs.geometry], {})


@pytest.mark.parametrize(("raw", "value"), [("1,234.50", 1234.5), ("(12)", -12), ("12%", 0.12), ("007", "007"),
                                            ("abc", "abc"), ("₹ 1,000", 1000)])
def test_to_number(raw, value):
    assert to_number(raw) == value


def test_excel_never_writes_formulas(tmp_path):
    import openpyxl

    doc = pymupdf.open()
    page = doc.new_page()
    from tests.fixtures.pdf.factory import _table

    _table(page, 50, 50, [["Name", "Value"], ["=cmd|' /C calc'!A0", "1"], ["@SUM(A1)", "2"]], col_w=160)
    path = tmp_path / "evil.pdf"
    doc.save(path)
    ctx = run_tool("pdf.to_excel", tmp_path, [path], {})
    ws = openpyxl.load_workbook(ctx.outputs[0][0]).active
    assert ws["A2"].data_type == "s" and ws["A3"].data_type == "s"


def test_pdf_to_markdown(pdfs, tmp_path):
    ctx = run_tool("pdf.to_markdown", tmp_path, [pdfs.report], {})
    md = ctx.outputs[0][0].read_text("utf-8")
    assert "# Chapter 1: Quarterly results" in md
    assert "| Region | Q1 | Q2 | Q3 |" in md and "|---|---|---|---|" in md
    assert "- Revenue up twelve percent" in md and "1. Approve budget" in md
    assert "[Visit our website](https://example.com/report)" in md
    assert "ACME Corp" not in md  # running header removed
    assert "*[Image]*" in md
    assert ctx.outputs[0][1] == "report.md" and ctx.data["tables"] == 3
    keep = run_tool("pdf.to_markdown", tmp_path, [pdfs.report], {"headers_footers": "keep", "images": "omit",
                                                                 "page_breaks": True})
    md2 = keep.outputs[0][0].read_text("utf-8")
    assert "ACME Corp" in md2 and "*[Image]*" not in md2 and "\n---\n" in md2


def test_markdown_escapes_markup_and_unsafe_links(tmp_path):
    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_text((72, 72), "# not a heading <script>alert(1)</script> *x*", fontsize=11)
    page.insert_text((72, 100), "click", fontsize=11)
    r = page.search_for("click")[0]
    page.insert_link({"kind": pymupdf.LINK_URI, "from": r, "uri": "javascript:alert(1)"})
    path = tmp_path / "md.pdf"
    doc.save(path)
    md = run_tool("pdf.to_markdown", tmp_path, [path], {}).outputs[0][0].read_text("utf-8")
    assert "\\#" in md and "<script>" not in md and "\\*x\\*" in md
    assert "javascript:" not in md


def test_markdown_warns_on_scans(pdfs, tmp_path):
    ctx = run_tool("pdf.to_markdown", tmp_path, [pdfs.scanned], {"pages": "1"})
    assert any("OCR" in w for w in ctx.warnings)


def test_outputs_are_valid_office_zips(pdfs, tmp_path):
    for tool in ("pdf.to_word", "pdf.to_excel"):
        ctx = run_tool(tool, tmp_path, [pdfs.report], {})
        with zipfile.ZipFile(io.BytesIO(ctx.outputs[0][0].read_bytes())) as zf:
            assert zf.testzip() is None
